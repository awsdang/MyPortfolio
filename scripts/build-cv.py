"""Build the editable CV from its Markdown source using python-docx.

Run with the Codex bundled Python, then render with the documents skill's
render_docx.py --emit_pdf. The PDF is visually reviewed before distribution.
"""
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "cv/Aws_Abdulfattah__Senior_Software_Engineer.md"
OUTPUT = SOURCE.with_suffix(".docx")

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin = section.bottom_margin = Mm(13)
section.left_margin = section.right_margin = Mm(15)
styles = doc.styles
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"]:
    style = styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    style.paragraph_format.widow_control = True

# Strip inherited Word template borders and theme styling from the title block.
for style in styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    for spacing in list(style.element.iter(qn("w:spacing"))):
        if spacing.getparent().tag == qn("w:rPr"):
            spacing.getparent().remove(spacing)

normal = styles["Normal"]
normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = 1.04
normal.paragraph_format.space_after = Pt(2.5)
styles["Title"].font.size = Pt(23)
styles["Title"].font.bold = True
styles["Title"].paragraph_format.space_after = Pt(2)
styles["Subtitle"].font.size = Pt(12)
styles["Subtitle"].font.bold = True
styles["Subtitle"].paragraph_format.space_after = Pt(2)
styles["Heading 1"].font.size = Pt(10.5)
styles["Heading 1"].font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(7)
styles["Heading 1"].paragraph_format.space_after = Pt(3)
styles["Heading 2"].font.size = Pt(10)
styles["Heading 2"].font.bold = True
styles["Heading 2"].paragraph_format.space_before = Pt(4)
styles["Heading 2"].paragraph_format.space_after = Pt(1)
bullet = styles["List Bullet"]
bullet.paragraph_format.left_indent = Mm(3)
bullet.paragraph_format.first_line_indent = Mm(-3)
bullet.paragraph_format.space_after = Pt(3)
bullet.paragraph_format.line_spacing = 1.04
# A literal Unicode bullet stays readable in PDF and plain-text extraction.
num_pr = OxmlElement("w:numPr")
num_id = OxmlElement("w:numId")
num_id.set(qn("w:val"), "0")
num_pr.append(num_id)
bullet.element.get_or_add_pPr().append(num_pr)

URLS = {
    "awsdang@gmail.com": "mailto:awsdang@gmail.com",
    "awsdang.com": "https://awsdang.com/",
    "linkedin.com/in/awsdang": "https://www.linkedin.com/in/awsdang",
    "github.com/awsdang": "https://github.com/awsdang",
    "github.com/awsdang/notifyx": "https://github.com/awsdang/notifyx",
}
url_re = re.compile("(" + "|".join(re.escape(s) for s in sorted(URLS, key=len, reverse=True)) + ")")

def text_runs(p, content):
    for n, piece in enumerate(re.split(r"\*\*(.*?)\*\*", content)):
        for token in url_re.split(piece):
            if not token:
                continue
            if token in URLS:
                h = OxmlElement("w:hyperlink")
                h.set(qn("r:id"), p.part.relate_to(URLS[token], RT.HYPERLINK, is_external=True))
                r = OxmlElement("w:r")
                props = OxmlElement("w:rPr")
                if n % 2:
                    props.append(OxmlElement("w:b"))
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "000000")
                props.append(color)
                r.append(props)
                t = OxmlElement("w:t")
                t.text = token
                r.append(t)
                h.append(r)
                p._p.append(h)
            else:
                run = p.add_run(token)
                if n % 2:
                    run.bold = True

in_header = True
after_role = False
for line in SOURCE.read_text().splitlines():
    if not line.strip():
        continue
    if line == "<!-- pagebreak -->":
        doc.add_page_break()
    elif line.startswith("# "):
        text_runs(doc.add_paragraph(style="Title"), line[2:])
    elif line.startswith("## "):
        in_header = False
        text_runs(doc.add_paragraph(style="Heading 1"), line[3:])
        after_role = False
    elif line.startswith("### "):
        text_runs(doc.add_paragraph(style="Heading 2"), line[4:])
        after_role = True
    elif line.startswith("- "):
        text_runs(doc.add_paragraph(style="List Bullet"), "• " + line[2:])
        after_role = False
    else:
        p = doc.add_paragraph(style="Subtitle" if in_header and line == "Senior Software Engineer" else "Normal")
        text_runs(p, line)
        if in_header:
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.keep_with_next = True
            if line != "Senior Software Engineer":
                for r in p.runs:
                    r.font.size = Pt(9.5)
        elif after_role:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(9)
            after_role = False
        elif line.startswith("**") and " | " in line:
            p.paragraph_format.keep_with_next = not line.startswith("**University of Technology")

doc.core_properties.title = "Aws Abdulfattah Senior Software Engineer"
doc.core_properties.subject = "Senior software engineering and technical leadership"
doc.core_properties.author = "Aws Abdulfattah"
doc.core_properties.keywords = "Senior Software Engineer, Technical Lead, Payments, TypeScript, React, PostgreSQL"
doc.core_properties.comments = ""
doc.save(OUTPUT)
print(OUTPUT)
print(f"Source words: {len(SOURCE.read_text().split())}")
